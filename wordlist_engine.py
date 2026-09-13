from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
from html import escape, unescape
from html.parser import HTMLParser
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterator
from urllib.parse import quote

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openai import OpenAI

COLUMN_WIDTHS_CM = (4.4, 2.2, 10.8)
TABLE_WIDTH_CM = sum(COLUMN_WIDTHS_CM)
SUPPORTED_INPUT_SUFFIXES = {".docx", ".txt"}
ProgressCallback = Callable[[int, str], None]
DEFAULT_HOMEPAGE_BASE_URL = "https://xxxcillian.org"
DEFAULT_PROVIDER_ID = "bigmodel"


@dataclass(frozen=True)
class ApiProvider:
    id: str
    display_name: str
    base_url: str
    default_models: tuple[str, ...]
    env_var: str
    keyring_username: str
    raw_result_label: str

    @property
    def default_model(self) -> str:
        return self.default_models[0]


API_PROVIDERS: dict[str, ApiProvider] = {
    "bigmodel": ApiProvider(
        id="bigmodel",
        display_name="BigModel / 智谱中国版",
        base_url="https://open.bigmodel.cn/api/paas/v4",
        default_models=("glm-5.3-flash", "glm-5.3"),
        env_var="BIGMODEL_API_KEY",
        keyring_username="bigmodel_api_key",
        raw_result_label="BigModel",
    ),
    "zai": ApiProvider(
        id="zai",
        display_name="Ox Alpha / Z.ai GLM",
        base_url="https://api.z.ai/api/paas/v4/",
        default_models=("glm-5.3-flash", "glm-5.3"),
        env_var="ZAI_API_KEY",
        keyring_username="zai_api_key",
        raw_result_label="OxAlpha",
    ),
    "moonshot": ApiProvider(
        id="moonshot",
        display_name="Kimi / Moonshot",
        base_url="https://api.moonshot.cn/v1",
        default_models=("kimi-k2.6", "kimi-k2.5"),
        env_var="MOONSHOT_API_KEY",
        keyring_username="moonshot_api_key",
        raw_result_label="Kimi",
    ),
}


def get_api_provider(provider_id: str | None) -> ApiProvider:
    return API_PROVIDERS.get(provider_id or "", API_PROVIDERS[DEFAULT_PROVIDER_ID])


@dataclass(frozen=True)
class HomepagePublishResult:
    archive_dir: Path
    index_path: Path
    commit_hash: str | None = None
    url: str | None = None


@dataclass(frozen=True)
class ProcessingResult:
    output_docx: Path
    output_html: Path
    raw_json: Path
    parsed_json: Path
    model: str
    homepage_publish: HomepagePublishResult | None = None
    prompt_tokens: int | None = None
    completion_tokens: int | None = None
    total_tokens: int | None = None


def emit(callback: ProgressCallback | None, percent: int, message: str) -> None:
    if callback:
        callback(percent, message)


def iter_block_items(parent: DocumentObject | _Cell) -> Iterator[Paragraph | Table]:
    """按 Word 文档中的实际顺序遍历段落和表格。"""
    if isinstance(parent, DocumentObject):
        parent_element = parent.element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    else:
        raise TypeError(f"不支持的 Word 容器类型：{type(parent)!r}")

    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def extract_docx_content(path: Path) -> str:
    document = Document(path)
    blocks: list[str] = []

    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            value = block.text.strip()
            if value:
                blocks.append(value)
            continue

        for row in block.rows:
            values: list[str] = []
            for cell in row.cells:
                cell_text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )
                values.append(cell_text)

            if any(values):
                blocks.append("【原表格行】" + " ｜ ".join(values))

    content = "\n".join(blocks).strip()
    if not content:
        raise ValueError("输入 Word 文档中没有提取到可处理的文字。")
    return content


def extract_txt_content(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "gb18030"):
        try:
            content = path.read_text(encoding=encoding).strip()
            break
        except UnicodeDecodeError:
            continue
    else:
        content = path.read_text(encoding="utf-8", errors="replace").strip()

    if not content:
        raise ValueError("输入 TXT 文档中没有提取到可处理的文字。")
    return content


def extract_input_content(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx_content(path)
    if suffix == ".txt":
        return extract_txt_content(path)
    raise FileNotFoundError("请选择一个有效的 .docx 或 .txt 原始单词记录文件。")


def build_messages(rules: str, source_text: str) -> list[dict[str, str]]:
    schema = r'''
{
  "document_title": "IELTS 单词本整理版",
  "subtitle": "词族成组 · 同义辨析 · 雅思实用频度逐词标注",
  "corrections": [
    {
      "original": "错误形式",
      "corrected": "正确形式",
      "note": "简短说明"
    }
  ],
  "sections": [
    {
      "title": "一、名词与名词核心词族",
      "entries": [
        {
          "group": "左栏显示的核心词条、词族或近义组",
          "frequency_range": "A-C",
          "items": [
            {
              "headword": "单词或短语",
              "pos": "n./v./adj./adv./phr. v. 等",
              "frequency": "A、B、C、D 或 A/B",
              "definition": "准确、自然的中文释义和必要语义限制",
              "collocations": ["常用搭配一", "常用搭配二"],
              "usage_note": "没有必要说明时写空字符串"
            }
          ],
          "family_additions": [
            {
              "headword": "必要的常用词族补充",
              "pos": "词性",
              "frequency": "A/B/C/D",
              "definition": "简短释义"
            }
          ],
          "difference": "同组词之间的 nuance、替换限制和语境区别",
          "group_note": "其他必要说明；没有则写空字符串"
        }
      ]
    },
    {"title": "二、动词与动词核心词族", "entries": []},
    {"title": "三、形容词与形容词核心词族", "entries": []},
    {"title": "四、副词与表达偏好", "entries": []}
  ]
}
'''.strip()

    system_message = (
        "你是一名严谨的英语词汇编辑和雅思词汇资料整理专家。"
        "必须严格执行用户提供的工作规范，并只返回合法 JSON 对象。"
        "不得返回 Markdown、代码围栏、解释性前言或结语。"
    )

    user_message = f"""
请整理下面的原始英文单词记录。

【必须执行的完整工作规范】
{rules}

【额外结构化输出要求】
1. 只能返回一个合法 JSON 对象，键名和层级严格采用下方结构。
2. sections 必须按名词、动词、形容词、副词／表达的顺序排列。
3. 词族完整性优先，不得为了分板块而拆散同一词族。
4. 原始记录中的每一个英文词、词组和固定搭配都必须至少出现在 group、items、family_additions、difference、group_note 或 corrections 之一。
5. 每个 items 项目必须单独标注雅思实用频度；frequency_range 只用于左栏概括。
6. collocations 通常每词保留 1—4 个真正常用、具有复习价值的搭配。
7. 不要为了“完整”添加大量冷僻派生词。
8. definition、usage_note、difference 和 group_note 使用简体中文。
9. corrections 只记录原始材料中真实存在的拼写、词形或记录错误；没有则返回空数组。
10. 请在输出前自行检查是否遗漏、错误合并、错误拆散或误标频度。

【必须采用的 JSON 结构】
{schema}

【原始英文单词记录开始】
{source_text}
【原始英文单词记录结束】
""".strip()

    return [
        {"role": "system", "content": system_message},
        {"role": "user", "content": user_message},
    ]


def parse_json_response(raw_text: str) -> dict[str, Any]:
    raw_text = raw_text.strip()
    try:
        data = json.loads(raw_text)
    except json.JSONDecodeError as exc:
        first = raw_text.find("{")
        last = raw_text.rfind("}")
        if first == -1 or last <= first:
            raise ValueError(f"模型返回内容不是合法 JSON：{exc}") from exc
        try:
            data = json.loads(raw_text[first : last + 1])
        except json.JSONDecodeError as second_exc:
            raise ValueError(f"模型返回的 JSON 无法解析：{second_exc}") from second_exc

    if not isinstance(data, dict):
        raise ValueError("模型返回的 JSON 顶层不是对象。")
    return data


def validate_data(data: dict[str, Any]) -> None:
    sections = data.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("整理结果缺少有效的 sections 列表。")

    entry_count = 0
    for section in sections:
        if not isinstance(section, dict):
            raise ValueError("sections 中存在非对象项目。")
        title = str(section.get("title", "")).strip()
        entries = section.get("entries", [])
        if not title:
            raise ValueError("存在没有标题的章节。")
        if not isinstance(entries, list):
            raise ValueError(f"章节“{title}”的 entries 不是列表。")

        for entry in entries:
            if not isinstance(entry, dict):
                raise ValueError(f"章节“{title}”中存在非对象词条。")
            if not str(entry.get("group", "")).strip():
                raise ValueError(f"章节“{title}”中存在缺少 group 的词条。")
            items = entry.get("items")
            if not isinstance(items, list) or not items:
                raise ValueError(f"词条“{entry.get('group')}”缺少 items。")
            entry_count += 1

    if entry_count == 0:
        raise ValueError("整理结果没有生成任何词条。")


def build_extra_body(provider: ApiProvider, thinking_enabled: bool) -> dict[str, Any]:
    if provider.id in {"bigmodel", "zai"}:
        return {
            "thinking": {
                "type": "enabled",
                "clear_thinking": False,
            },
            "reasoning_effort": "max" if thinking_enabled else "low",
        }
    return {
        "thinking": {
            "type": "enabled" if thinking_enabled else "disabled"
        }
    }


def call_model(
    *,
    provider_id: str,
    api_key: str,
    model: str,
    rules: str,
    source_text: str,
    thinking_enabled: bool,
) -> tuple[str, Any]:
    provider = get_api_provider(provider_id)
    client = OpenAI(
        api_key=api_key,
        base_url=provider.base_url,
        timeout=900.0,
        max_retries=2,
    )

    request_args: dict[str, Any] = {
        "model": model,
        "messages": build_messages(rules, source_text),
        "response_format": {"type": "json_object"},
        "extra_body": build_extra_body(provider, thinking_enabled),
    }
    if provider.id in {"bigmodel", "zai"}:
        request_args["max_tokens"] = 65536
    else:
        request_args["max_completion_tokens"] = 65536

    response = client.chat.completions.create(**request_args)

    if not response.choices:
        raise ValueError(f"{provider.display_name} 没有返回任何候选结果。")

    choice = response.choices[0]
    if choice.finish_reason == "length":
        raise ValueError(f"{provider.display_name} 输出达到长度上限，JSON 可能被截断。请减少输入内容后重试。")

    content = choice.message.content
    if not content:
        raise ValueError(f"{provider.display_name} 返回了空内容。")
    return content.strip(), response.usage


def value_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_run_font(run, *, size: float = 9, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_format(paragraph, *, space_after: float = 3, line_spacing: float = 1.08) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def set_fixed_table_layout(table, total_width_cm: float) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(Cm(total_width_cm).twips))
    width.set(qn("w:type"), "dxa")


def set_table_grid_widths(table, widths_cm: tuple[float, ...]) -> None:
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width_cm in widths_cm:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(Cm(width_cm).twips))
        grid.append(column)


def set_cell_width(cell, width_cm: float) -> None:
    width = Cm(width_cm)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(width.twips))
    tc_width.set(qn("w:type"), "dxa")


def apply_column_widths(cells) -> None:
    for cell, width_cm in zip(cells, COLUMN_WIDTHS_CM):
        set_cell_width(cell, width_cm)


def format_item_detail(item: dict[str, Any]) -> str:
    parts: list[str] = []
    definition = value_text(item.get("definition"))
    if definition:
        parts.append(definition.rstrip("。"))

    collocations = item.get("collocations", [])
    if isinstance(collocations, list):
        clean = [value_text(value) for value in collocations if value_text(value)]
        if clean:
            parts.append(" / ".join(clean).rstrip("。"))

    usage_note = value_text(item.get("usage_note"))
    if usage_note:
        parts.append(f"【使用提醒】{usage_note.rstrip('。')}")

    return "。".join(parts) + ("。" if parts else "")


def add_labeled_paragraph(cell, label: str, content: str) -> None:
    if not content:
        return
    paragraph = cell.add_paragraph()
    set_paragraph_format(paragraph, space_after=4)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=9, bold=True, color="1F4E78")
    content_run = paragraph.add_run(content)
    set_run_font(content_run, size=9)


def add_item_to_cell(cell, item: dict[str, Any], *, first: bool) -> None:
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    set_paragraph_format(paragraph, space_after=5)

    headword = value_text(item.get("headword")) or "未命名词条"
    pos = value_text(item.get("pos"))
    frequency = value_text(item.get("frequency"))

    label = headword
    if pos:
        label += f" ({pos})"
    if frequency:
        label += f"｜{frequency}"
    label += "："

    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=9, bold=True)
    detail_run = paragraph.add_run(format_item_detail(item))
    set_run_font(detail_run, size=9)


def format_family_additions(values: Any) -> str:
    if not isinstance(values, list):
        return ""
    parts: list[str] = []
    for item in values:
        if not isinstance(item, dict):
            continue
        headword = value_text(item.get("headword"))
        if not headword:
            continue
        label = headword
        pos = value_text(item.get("pos"))
        frequency = value_text(item.get("frequency"))
        definition = value_text(item.get("definition"))
        if pos:
            label += f" ({pos})"
        if frequency:
            label += f"｜{frequency}"
        if definition:
            label += f"：{definition}"
        parts.append(label)
    return "；".join(parts)


def add_word_row(table, entry: dict[str, Any]) -> None:
    cells = table.add_row().cells
    apply_column_widths(cells)
    for cell in cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        set_cell_margins(cell)

    group_p = cells[0].paragraphs[0]
    set_paragraph_format(group_p)
    group_run = group_p.add_run(value_text(entry.get("group")))
    set_run_font(group_run, size=9, bold=True)

    set_cell_shading(cells[1], "E2F0D9")
    freq_p = cells[1].paragraphs[0]
    freq_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(freq_p)
    freq_run = freq_p.add_run(value_text(entry.get("frequency_range")) or "逐词")
    set_run_font(freq_run, size=9, bold=True, color="1F4E78")

    note_p = cells[1].add_paragraph()
    note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(note_p)
    note_run = note_p.add_run("逐词见右栏")
    set_run_font(note_run, size=8, color="595959")

    items = entry.get("items", [])
    for index, item in enumerate(items):
        if isinstance(item, dict):
            add_item_to_cell(cells[2], item, first=index == 0)

    add_labeled_paragraph(cells[2], "【词族补充】：", format_family_additions(entry.get("family_additions")))
    add_labeled_paragraph(cells[2], "【区别】：", value_text(entry.get("difference")))
    add_labeled_paragraph(cells[2], "【补充说明】：", value_text(entry.get("group_note")))


def add_table_header(table) -> None:
    cells = table.rows[0].cells
    apply_column_widths(cells)
    set_repeat_table_header(table.rows[0])
    titles = ["词条 / 词族", "雅思频度", "中文释义、nuance 与搭配"]
    for cell, title in zip(cells, titles):
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=80, bottom=80)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(paragraph, space_after=0)
        run = paragraph.add_run(title)
        set_run_font(run, size=9, bold=True, color="FFFFFF")


def add_info_box(document) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_fixed_table_layout(table, TABLE_WIDTH_CM)
    set_table_grid_widths(table, (TABLE_WIDTH_CM,))

    cell = table.cell(0, 0)
    set_cell_width(cell, TABLE_WIDTH_CM)
    set_cell_shading(cell, "DDEBF7")
    set_cell_margins(cell, top=110, bottom=110, start=130, end=130)

    lines = [
        "本版标准：按雅思备考实用性逐词标注频度，不按词长或生僻程度机械分级。",
        "A｜高频必掌握：四科常见，或写作、口语改写价值高；要求会认、会搭配并能主动使用。",
        "B｜常见且值得掌握：特定话题、正式表达或学术语境有价值；熟悉主要用法并争取主动使用。",
        "C｜低频，认读为主：可能出现在阅读或窄语境中，但主动产出收益较低。",
        "D｜极低频/不建议主动使用：过时、生硬、文学化或专业范围过窄；优先采用更自然替代表达。",
        "词族原则：同一词族、近义组、反义词和形近词集中对照；关系不等同的词明确区分。",
    ]

    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        set_paragraph_format(paragraph, space_after=2)
        run = paragraph.add_run(line)
        set_run_font(run, size=8.5, bold=index == 0)


def add_corrections(document, corrections: Any) -> None:
    if not isinstance(corrections, list):
        return
    parts: list[str] = []
    for item in corrections:
        if not isinstance(item, dict):
            continue
        original = value_text(item.get("original"))
        corrected = value_text(item.get("corrected"))
        note = value_text(item.get("note"))
        if not original and not corrected:
            continue
        part = f"{original} → {corrected}"
        if note:
            part += f"（{note}）"
        parts.append(part)

    if not parts:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    title_run = paragraph.add_run("本次原始记录中的主要纠错：")
    set_run_font(title_run, size=9, bold=True, color="1F4E78")
    body_run = paragraph.add_run("；".join(parts) + "。")
    set_run_font(body_run, size=9)


def sanitize_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*]+', "_", value).strip(" .")
    return value or "单词本"


def sanitize_url_part(value: str) -> str:
    value = value.strip().lower()
    value = re.sub(r"\s+", "-", value)
    value = re.sub(r"[^\w\-\u4e00-\u9fff]+", "-", value, flags=re.UNICODE)
    value = re.sub(r"-+", "-", value).strip("-_")
    return value[:80] or "wordbook"


def generate_word(
    data: dict[str, Any],
    source_path: Path,
    output_dir: Path,
    *,
    stamp: str | None = None,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now()
    date_text = now.strftime("%Y-%m-%d")
    stamp = stamp or now.strftime("%Y%m%d_%H%M%S")
    stem = sanitize_filename(source_path.stem)
    output_path = output_dir / f"{stem}_雅思单词本_整理版_{stamp}.docx"

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    title_run = title.add_run(f"IELTS 单词本 · {date_text}")
    set_run_font(title_run, size=16, bold=True, color="1F4E78")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle_run = subtitle.add_run(
        value_text(data.get("subtitle")) or "词族成组 · 同义辨析 · 雅思实用频度逐词标注"
    )
    set_run_font(subtitle_run, size=9, color="595959")

    add_info_box(document)
    add_corrections(document, data.get("corrections"))

    for section_data in data.get("sections", []):
        if not isinstance(section_data, dict):
            continue
        entries = section_data.get("entries", [])
        if not isinstance(entries, list) or not entries:
            continue

        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(5)
        heading_run = heading.add_run(value_text(section_data.get("title")))
        set_run_font(heading_run, size=12, bold=True, color="1F4E78")

        table = document.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        set_fixed_table_layout(table, TABLE_WIDTH_CM)
        set_table_grid_widths(table, COLUMN_WIDTHS_CM)
        add_table_header(table)

        for entry in entries:
            if isinstance(entry, dict):
                add_word_row(table, entry)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(f"IELTS 单词本 · {date_text}")
    set_run_font(footer_run, size=8, color="808080")

    document.save(output_path)
    return output_path


def html_value(value: Any) -> str:
    return escape(value_text(value), quote=True)


def split_reading_chunks(value: str) -> list[str]:
    value = value_text(value)
    if not value:
        return []
    if " / " not in value or len(value) < 90:
        return [value]
    chunks = [chunk.strip() for chunk in re.split(r"\s+/\s+", value) if chunk.strip()]
    return chunks or [value]


def split_inline_note(value: str) -> tuple[str, str]:
    match = re.search(r"(。|；)?(【(?:使用提醒|区别|补充说明)】[：:]?.*)", value)
    if not match:
        return value.strip(), ""
    main = value[: match.start(2)].strip()
    note = match.group(2).strip()
    return main, note


def split_definition_phrase(value: str) -> tuple[str, str]:
    match = re.search(r"。(?=\s*[A-Za-z])", value)
    if not match:
        return value.strip(), ""
    definition = value[: match.end()].strip()
    phrase = value[match.end() :].strip()
    return definition, phrase


def is_word_definition(value: str) -> bool:
    return bool(re.match(r"^[A-Za-z][^。]{0,120}(?:\([^)]*\)|[｜|]\s*[A-D])[^。]*[：:]", value))


def render_phrase_card(phrases: list[str]) -> str:
    clean = [html_value(value) for value in phrases if value_text(value)]
    if not clean:
        return ""
    items = "".join(f"<li>{value}</li>" for value in clean)
    return f'<div class="phrase-card"><h4>短语 / 搭配</h4><ul class="phrase-list">{items}</ul></div>'


def render_note_card(value: str) -> str:
    class_name = "note-point reminder-point" if value.startswith("【使用提醒】") else "note-point label-point"
    return f'<p class="{class_name}">{html_value(value)}</p>'


def render_reading_blocks(value: Any) -> str:
    chunks = split_reading_chunks(value_text(value))
    if not chunks:
        return ""
    if len(chunks) == 1:
        return f'<p class="definition">{html_value(chunks[0])}</p>'

    blocks: list[str] = []
    phrase_group: list[str] = []
    pending_note_index: int | None = None
    pending_note_value = ""

    def flush_phrases() -> None:
        if phrase_group:
            blocks.append(render_phrase_card(phrase_group))
            phrase_group.clear()

    def append_note(value: str) -> None:
        nonlocal pending_note_index, pending_note_value
        flush_phrases()
        pending_note_value = value
        blocks.append(render_note_card(value))
        pending_note_index = len(blocks) - 1

    def extend_pending_note(value: str) -> bool:
        nonlocal pending_note_value
        if pending_note_index is None:
            return False
        pending_note_value = f"{pending_note_value} / {value}"
        blocks[pending_note_index] = render_note_card(pending_note_value)
        return True

    for chunk in chunks:
        main, note = split_inline_note(chunk)
        if main:
            if is_word_definition(main):
                pending_note_index = None
                flush_phrases()
                definition, phrase = split_definition_phrase(main)
                if definition:
                    blocks.append(f'<p class="detail-point definition-point">{html_value(definition)}</p>')
                if phrase:
                    phrase_group.append(phrase)
            elif main.startswith("【"):
                append_note(main)
            elif extend_pending_note(main):
                continue
            else:
                pending_note_index = None
                phrase_group.append(main)

        if note:
            append_note(note)

    flush_phrases()
    return '<div class="detail-list">' + "".join(blocks) + "</div>"


def render_mobile_item(item: dict[str, Any], *, parent_group: str = "") -> str:
    raw_headword = value_text(item.get("headword")) or "未命名词条"
    headword = html_value(raw_headword)
    pos = html_value(item.get("pos"))
    frequency = html_value(item.get("frequency"))
    usage_note = html_value(item.get("usage_note"))
    show_header = raw_headword != value_text(parent_group)

    meta_parts = [part for part in (pos, frequency) if part]
    meta = f'<span class="meta">{" · ".join(meta_parts)}</span>' if meta_parts else ""
    header_html = (
        f"""
          <header>
            <h3>{headword}</h3>
            {meta}
          </header>
        """
        if show_header
        else ""
    )

    collocations_html = ""
    collocations = item.get("collocations", [])
    if isinstance(collocations, list):
        clean = [html_value(value) for value in collocations if value_text(value)]
        if clean:
            collocations_html = (
                '<div class="chips">'
                + "".join(f"<span>{value}</span>" for value in clean)
                + "</div>"
            )

    usage_html = f'<p class="usage">使用提醒：{usage_note}</p>' if usage_note else ""
    definition_html = render_reading_blocks(item.get("definition"))

    return f"""
        <article class="word">
          {header_html}
          {definition_html}
          {collocations_html}
          {usage_html}
        </article>
    """


def render_family_additions(values: Any) -> str:
    if not isinstance(values, list):
        return ""

    rows: list[str] = []
    for item in values:
        if not isinstance(item, dict):
            continue
        headword = html_value(item.get("headword"))
        if not headword:
            continue
        meta_parts = [html_value(item.get("pos")), html_value(item.get("frequency"))]
        meta = " · ".join(part for part in meta_parts if part)
        definition = html_value(item.get("definition"))
        rows.append(
            f"<li><strong>{headword}</strong>"
            f"{f' <span>{meta}</span>' if meta else ''}"
            f"{f'：{definition}' if definition else ''}</li>"
        )

    if not rows:
        return ""
    return '<div class="note-block"><h4>词族补充</h4><ul>' + "".join(rows) + "</ul></div>"


def render_mobile_entry(entry: dict[str, Any]) -> str:
    raw_group = value_text(entry.get("group"))
    group = html_value(raw_group)
    frequency_range = html_value(entry.get("frequency_range")) or "逐词"

    items_html: list[str] = []
    items = entry.get("items", [])
    if isinstance(items, list):
        for item in items:
            if isinstance(item, dict):
                items_html.append(render_mobile_item(item, parent_group=raw_group))

    difference = html_value(entry.get("difference"))
    group_note = html_value(entry.get("group_note"))
    difference_html = (
        f'<div class="note-block"><h4>区别</h4><p>{difference}</p></div>'
        if difference
        else ""
    )
    note_html = (
        f'<div class="note-block"><h4>补充说明</h4><p>{group_note}</p></div>'
        if group_note
        else ""
    )

    return f"""
      <article class="entry" data-entry>
        <div class="entry-title">
          <h2>{group}</h2>
          <span>{frequency_range}</span>
        </div>
        <div class="words">
          {"".join(items_html)}
        </div>
        {render_family_additions(entry.get("family_additions"))}
        {difference_html}
        {note_html}
      </article>
    """


def render_corrections(corrections: Any) -> str:
    if not isinstance(corrections, list):
        return ""

    rows: list[str] = []
    for item in corrections:
        if not isinstance(item, dict):
            continue
        original = html_value(item.get("original"))
        corrected = html_value(item.get("corrected"))
        note = html_value(item.get("note"))
        if not original and not corrected:
            continue
        rows.append(
            f"<li><strong>{original}</strong> → <strong>{corrected}</strong>"
            f"{f'（{note}）' if note else ''}</li>"
        )

    if not rows:
        return ""
    return """
      <section class="corrections">
        <h2>本次纠错</h2>
        <ul>
          %s
        </ul>
      </section>
    """ % "".join(rows)


def generate_mobile_html(
    data: dict[str, Any],
    source_path: Path,
    output_dir: Path,
    *,
    stamp: str | None = None,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now()
    date_text = now.strftime("%Y-%m-%d")
    stamp = stamp or now.strftime("%Y%m%d_%H%M%S")
    stem = sanitize_filename(source_path.stem)
    output_path = output_dir / f"{stem}_雅思单词本_手机版_{stamp}.html"

    sections_html: list[str] = []
    total_entries = 0
    total_words = 0
    for index, section_data in enumerate(data.get("sections", [])):
        if not isinstance(section_data, dict):
            continue
        entries = section_data.get("entries", [])
        if not isinstance(entries, list) or not entries:
            continue

        rendered_entries: list[str] = []
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            rendered_entries.append(render_mobile_entry(entry))
            total_entries += 1
            items = entry.get("items", [])
            if isinstance(items, list):
                total_words += sum(1 for item in items if isinstance(item, dict))

        if not rendered_entries:
            continue

        open_attr = " open" if index == 0 else ""
        sections_html.append(
            f"""
            <details class="section"{open_attr}>
              <summary>{html_value(section_data.get("title"))}</summary>
              {"".join(rendered_entries)}
            </details>
            """
        )

    subtitle = html_value(data.get("subtitle")) or "词族成组 · 同义辨析 · 雅思实用频度逐词标注"
    source_name = html_value(source_path.name)
    page_title = f"IELTS 单词本 · {date_text}"

    html = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
  <title>{escape(page_title, quote=True)}</title>
  <link rel="icon" href="data:,">
  <style>
    :root {{
      color-scheme: light dark;
      --bg: #f6f7f3;
      --panel: #ffffff;
      --ink: #18212b;
      --muted: #69727d;
      --line: #d9ded6;
      --accent: #1f6f8b;
      --accent-soft: #e3f3f8;
      --mark: #fff0bf;
      --shadow: 0 8px 26px rgba(22, 31, 41, 0.08);
    }}
    * {{ box-sizing: border-box; }}
    html {{ scroll-behavior: smooth; }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--ink);
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", "PingFang SC", "Microsoft YaHei", sans-serif;
      line-height: 1.62;
    }}
    body.dark {{
      --bg: #17191c;
      --panel: #22262b;
      --ink: #f2f4f7;
      --muted: #a8b0b8;
      --line: #3a4048;
      --accent: #7cc9de;
      --accent-soft: #223941;
      --mark: #594a1d;
      --shadow: none;
    }}
    .app {{
      width: min(100%, 760px);
      margin: 0 auto;
      padding: max(20px, env(safe-area-inset-top)) 14px max(32px, env(safe-area-inset-bottom));
    }}
    .hero {{
      padding: 10px 2px 14px;
    }}
    h1 {{
      margin: 0;
      font-size: 28px;
      line-height: 1.18;
      font-weight: 800;
    }}
    .subtitle {{
      margin: 8px 0 0;
      color: var(--muted);
      font-size: 15px;
    }}
    .stats {{
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 8px;
      margin: 16px 0;
    }}
    .stat, .toolbar, .corrections, .entry {{
      background: var(--panel);
      border: 1px solid var(--line);
      box-shadow: var(--shadow);
    }}
    .stat {{
      min-height: 70px;
      padding: 10px;
      border-radius: 8px;
    }}
    .stat strong {{
      display: block;
      font-size: 18px;
      line-height: 1.2;
    }}
    .stat span {{
      display: block;
      margin-top: 4px;
      color: var(--muted);
      font-size: 12px;
    }}
    .toolbar {{
      position: sticky;
      top: 0;
      z-index: 2;
      display: grid;
      grid-template-columns: 1fr auto;
      gap: 8px;
      padding: 10px;
      border-radius: 8px;
      margin: 0 0 14px;
    }}
    input {{
      width: 100%;
      min-height: 42px;
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 0 12px;
      background: var(--bg);
      color: var(--ink);
      font: inherit;
      outline: none;
    }}
    input:focus {{
      border-color: var(--accent);
      box-shadow: 0 0 0 3px color-mix(in srgb, var(--accent) 20%, transparent);
    }}
    button {{
      min-width: 46px;
      min-height: 42px;
      border: 1px solid var(--line);
      border-radius: 8px;
      background: var(--accent-soft);
      color: var(--ink);
      font: inherit;
      font-weight: 700;
    }}
    .hint {{
      grid-column: 1 / -1;
      color: var(--muted);
      font-size: 12px;
    }}
    .corrections {{
      border-radius: 8px;
      padding: 14px;
      margin-bottom: 14px;
    }}
    .corrections h2 {{
      margin: 0 0 8px;
      font-size: 17px;
    }}
    .corrections ul, .note-block ul {{
      margin: 0;
      padding-left: 18px;
    }}
    .section {{
      margin: 12px 0;
    }}
    summary {{
      min-height: 48px;
      display: flex;
      align-items: center;
      padding: 0 2px;
      color: var(--accent);
      font-size: 19px;
      font-weight: 800;
      cursor: pointer;
    }}
    .entry {{
      border-radius: 8px;
      padding: 16px;
      margin: 14px 0;
    }}
    .entry[hidden] {{
      display: none;
    }}
    .entry-title {{
      display: grid;
      grid-template-columns: 1fr;
      gap: 8px;
      align-items: start;
      border-bottom: 1px solid var(--line);
      padding-bottom: 12px;
      margin-bottom: 14px;
    }}
    .entry-title h2 {{
      margin: 0;
      font-size: 21px;
      line-height: 1.28;
      overflow-wrap: anywhere;
    }}
    .entry-title span {{
      display: inline-flex;
      align-items: center;
      justify-content: center;
      justify-self: start;
      min-width: 44px;
      min-height: 32px;
      border-radius: 999px;
      background: var(--accent-soft);
      color: var(--accent);
      font-size: 13px;
      font-weight: 800;
      white-space: nowrap;
    }}
    .word {{
      padding: 12px 0;
      border-bottom: 1px solid var(--line);
    }}
    .word:last-child {{
      border-bottom: 0;
    }}
    .word header {{
      display: flex;
      align-items: baseline;
      gap: 8px;
      flex-wrap: wrap;
    }}
    .word h3 {{
      margin: 0;
      font-size: 19px;
      line-height: 1.26;
      overflow-wrap: anywhere;
    }}
    .meta {{
      color: var(--accent);
      font-size: 13px;
      font-weight: 700;
    }}
    p {{
      margin: 7px 0 0;
    }}
    .definition {{
      font-size: 16px;
      line-height: 1.78;
      overflow-wrap: anywhere;
    }}
    .detail-list {{
      display: grid;
      gap: 10px;
      margin-top: 4px;
    }}
    .detail-point, .phrase-card, .note-point {{
      margin: 0;
      padding: 11px 12px;
      border: 1px solid color-mix(in srgb, var(--line) 78%, transparent);
      border-radius: 8px;
      background: color-mix(in srgb, var(--bg) 72%, var(--panel));
      font-size: 16px;
      line-height: 1.72;
      overflow-wrap: anywhere;
    }}
    .definition-point {{
      background: color-mix(in srgb, var(--panel) 82%, var(--bg));
    }}
    .phrase-card {{
      background: color-mix(in srgb, var(--accent-soft) 45%, var(--panel));
    }}
    .phrase-card h4 {{
      margin: 0 0 6px;
      color: var(--accent);
      font-size: 13px;
    }}
    .phrase-list {{
      display: grid;
      gap: 5px;
      margin: 0;
      padding-left: 20px;
    }}
    .label-point {{
      border-left: 4px solid var(--accent);
      background: var(--accent-soft);
    }}
    .reminder-point {{
      border-left: 4px solid #d39b2f;
      background: color-mix(in srgb, #fff0bf 42%, var(--panel));
    }}
    .chips {{
      display: flex;
      flex-wrap: wrap;
      gap: 6px;
      margin-top: 9px;
    }}
    .chips span {{
      border: 1px solid var(--line);
      border-radius: 999px;
      padding: 3px 9px;
      background: var(--bg);
      font-size: 13px;
    }}
    .usage {{
      color: var(--muted);
      font-size: 14px;
    }}
    .note-block {{
      margin-top: 12px;
      padding: 10px;
      border-left: 4px solid var(--accent);
      background: var(--accent-soft);
      border-radius: 6px;
    }}
    .note-block h4 {{
      margin: 0 0 5px;
      font-size: 14px;
    }}
    mark {{
      background: var(--mark);
      color: inherit;
      padding: 0 2px;
      border-radius: 3px;
    }}
    .empty {{
      display: none;
      padding: 20px 4px;
      color: var(--muted);
      text-align: center;
    }}
    .empty.show {{
      display: block;
    }}
    @media (min-width: 720px) {{
      .app {{ padding-left: 24px; padding-right: 24px; }}
      h1 {{ font-size: 34px; }}
      .entry {{ padding: 18px; }}
    }}
    @media (max-width: 520px) {{
      .stats {{ grid-template-columns: repeat(2, 1fr); }}
      .stat:last-child {{ grid-column: 1 / -1; }}
    }}
  </style>
</head>
<body>
  <main class="app">
    <header class="hero">
      <h1>{escape(page_title, quote=True)}</h1>
      <p class="subtitle">{subtitle}</p>
      <section class="stats" aria-label="概览">
        <div class="stat"><strong>{total_entries}</strong><span>词条/词族</span></div>
        <div class="stat"><strong>{total_words}</strong><span>逐词项目</span></div>
        <div class="stat"><strong>{date_text}</strong><span>生成日期</span></div>
      </section>
    </header>

    <section class="toolbar">
      <input id="search" type="search" placeholder="搜索单词、释义、搭配" autocomplete="off">
      <button id="theme" type="button" aria-label="切换深色模式">月</button>
      <div class="hint">来源：{source_name}。这是离线 HTML 文件，可在 iPhone Safari 或文件 App 中打开。</div>
    </section>

    {render_corrections(data.get("corrections"))}
    <div id="empty" class="empty">没有找到匹配内容。</div>
    {"".join(sections_html)}
  </main>

  <script>
    const search = document.getElementById("search");
    const empty = document.getElementById("empty");
    const entries = Array.from(document.querySelectorAll("[data-entry]"));
    const theme = document.getElementById("theme");

    function clearMarks(node) {{
      node.querySelectorAll("mark").forEach((mark) => mark.replaceWith(document.createTextNode(mark.textContent)));
      node.normalize();
    }}

    function markText(node, query) {{
      if (!query) return;
      const walker = document.createTreeWalker(node, NodeFilter.SHOW_TEXT, {{
        acceptNode(textNode) {{
          const parent = textNode.parentElement;
          if (!parent || parent.tagName === "SCRIPT" || parent.tagName === "STYLE" || parent.tagName === "MARK") {{
            return NodeFilter.FILTER_REJECT;
          }}
          return textNode.nodeValue.toLowerCase().includes(query)
            ? NodeFilter.FILTER_ACCEPT
            : NodeFilter.FILTER_SKIP;
        }}
      }});
      const nodes = [];
      while (walker.nextNode()) nodes.push(walker.currentNode);
      nodes.forEach((textNode) => {{
        const text = textNode.nodeValue;
        const lower = text.toLowerCase();
        const frag = document.createDocumentFragment();
        let start = 0;
        let index = lower.indexOf(query);
        while (index !== -1) {{
          frag.append(document.createTextNode(text.slice(start, index)));
          const mark = document.createElement("mark");
          mark.textContent = text.slice(index, index + query.length);
          frag.append(mark);
          start = index + query.length;
          index = lower.indexOf(query, start);
        }}
        frag.append(document.createTextNode(text.slice(start)));
        textNode.replaceWith(frag);
      }});
    }}

    function applySearch() {{
      const query = search.value.trim().toLowerCase();
      let visible = 0;
      entries.forEach((entry) => {{
        clearMarks(entry);
        const match = !query || entry.textContent.toLowerCase().includes(query);
        entry.hidden = !match;
        if (match) {{
          visible += 1;
          entry.closest("details").open = true;
          markText(entry, query);
        }}
      }});
      empty.classList.toggle("show", visible === 0);
    }}

    search.addEventListener("input", applySearch);
    const initialQuery = new URLSearchParams(window.location.search).get("q");
    if (initialQuery) {{
      search.value = initialQuery;
      applySearch();
      const firstMatch = entries.find((entry) => !entry.hidden);
      firstMatch?.scrollIntoView({{ behavior: "smooth", block: "start" }});
    }}
    theme.addEventListener("click", () => {{
      document.body.classList.toggle("dark");
      localStorage.setItem("wordlist-theme", document.body.classList.contains("dark") ? "dark" : "light");
    }});
    if (localStorage.getItem("wordlist-theme") === "dark") {{
      document.body.classList.add("dark");
    }}
  </script>
</body>
</html>
"""

    output_path.write_text(html, encoding="utf-8")
    return output_path


def count_wordbook_data(data: dict[str, Any]) -> tuple[int, int]:
    entry_count = 0
    word_count = 0
    for section_data in data.get("sections", []):
        if not isinstance(section_data, dict):
            continue
        entries = section_data.get("entries", [])
        if not isinstance(entries, list):
            continue
        for entry in entries:
            if not isinstance(entry, dict):
                continue
            entry_count += 1
            items = entry.get("items", [])
            if isinstance(items, list):
                word_count += sum(1 for item in items if isinstance(item, dict))
    return entry_count, word_count


def run_git(repo_path: Path, args: list[str]) -> str:
    completed = subprocess.run(
        ["git", "-C", str(repo_path), *args],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if completed.returncode != 0:
        details = (completed.stderr or completed.stdout).strip()
        raise ValueError(f"Git 命令失败：git {' '.join(args)}\n{details}")
    return completed.stdout.strip()


def read_git_short_head(repo_path: Path) -> str | None:
    for args in (["rev-parse", "--short", "HEAD"], ["log", "-1", "--format=%h"]):
        try:
            short_hash = run_git(repo_path, args)
        except ValueError:
            continue
        if short_hash:
            return short_hash
    return None


def validate_homepage_repo(repo_path: Path) -> None:
    if not repo_path.is_dir():
        raise FileNotFoundError(f"找不到个人主页仓库目录：{repo_path}")
    if not (repo_path / ".git").exists():
        raise ValueError(f"这不是 Git 仓库目录：{repo_path}")
    if not (repo_path / "notes" / "ielts").is_dir():
        raise ValueError("个人主页仓库中缺少 notes/ielts 雅思板块。")
    inside = run_git(repo_path, ["rev-parse", "--is-inside-work-tree"])
    if inside.strip().lower() != "true":
        raise ValueError(f"这不是有效的 Git 工作区：{repo_path}")


def build_wordbook_index_template(archive_html: str) -> str:
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>单词本收录 | 雅思</title>
  <meta name="description" content="归档 Wgen 生成的手机阅读版 HTML。" />
  <link rel="icon" href="data:," />
  <link rel="stylesheet" href="../../../style.css" />
</head>

<body>
  <header class="site-header">
    <nav class="nav">
      <a class="logo" href="../../../">裴天庆</a>
      <div class="nav-links">
        <a href="../../../">首页</a>
        <a href="../../">笔记</a>
        <a href="../">雅思</a>
        <a href="../../../#contact">联系</a>
      </div>
    </nav>
  </header>

  <main>
    <section class="notes-hero">
      <div class="breadcrumb">
        <a href="../../../">首页</a>
        <span>/</span>
        <a href="../../">学习笔记</a>
        <span>/</span>
        <a href="../">雅思</a>
        <span>/</span>
        <span>单词本收录</span>
      </div>
      <p class="eyebrow">IELTS · Wordbooks</p>
      <h1>单词本收录</h1>
      <p class="subtitle">这里会自动归档 Wgen 生成的手机阅读版 HTML。</p>
    </section>

    <section class="section wordbook-search-section">
      <h2>全库搜索</h2>
      <div class="wordbook-search-box">
        <label class="sr-only" for="wordbook-search">搜索归档单词本</label>
        <div class="wordbook-search-row">
          <input id="wordbook-search" type="search" placeholder="搜索单词、释义、搭配" autocomplete="off" />
          <button id="wordbook-search-clear" type="button" aria-label="清空搜索">×</button>
        </div>
        <p id="wordbook-search-status" class="wordbook-search-status">已索引归档词条，输入关键词即可检索。</p>
      </div>
      <div id="wordbook-search-results" class="wordbook-search-results" aria-live="polite"></div>
    </section>

    <section class="section">
      <h2>归档列表</h2>
      <!-- WORDBOOK_ARCHIVE_START -->
      {archive_html}
      <!-- WORDBOOK_ARCHIVE_END -->
    </section>
  </main>

  <footer class="footer">
    <p>© <span id="year"></span> 裴天庆. Built with curiosity.</p>
  </footer>

  <script src="../../../script.js"></script>
  <script src="wordbook-search.js"></script>
</body>
</html>
"""


def read_wordbook_metadata(wordbook_dir: Path) -> dict[str, Any] | None:
    meta_path = wordbook_dir / "meta.json"
    index_path = wordbook_dir / "index.html"
    if not meta_path.is_file() or not index_path.is_file():
        return None
    try:
        data = json.loads(meta_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    if not isinstance(data, dict):
        return None
    data["folder"] = wordbook_dir.name
    return data


def build_wordbook_archive_cards(wordbooks_dir: Path) -> str:
    items: list[dict[str, Any]] = []
    if wordbooks_dir.is_dir():
        for child in wordbooks_dir.iterdir():
            if child.is_dir():
                metadata = read_wordbook_metadata(child)
                if metadata:
                    items.append(metadata)

    items.sort(key=lambda item: value_text(item.get("generated_at")), reverse=True)
    if not items:
        return '<div class="note-slot">\n        <p>暂无单词本。整理器完成第一次网站归档后，会自动更新这里。</p>\n      </div>'

    cards: list[str] = []
    for item in items:
        folder = value_text(item.get("folder"))
        title = html_value(item.get("title")) or "雅思单词本"
        source_file = html_value(item.get("source_file"))
        generated_date = html_value(item.get("generated_date"))
        entry_count = html_value(item.get("entry_count"))
        word_count = html_value(item.get("word_count"))
        href = quote(folder) + "/"
        description_parts = []
        if source_file:
            description_parts.append(f"来源：{source_file}")
        if generated_date:
            description_parts.append(f"生成日期：{generated_date}")
        description = " · ".join(description_parts)
        cards.append(
            f"""<a class="archive-card" href="{href}">
        <div>
          <h3>{title}</h3>
          <p>{description}</p>
        </div>
        <div class="archive-meta">
          <span class="tag">{entry_count or "0"} 个词条/词族</span>
          <span class="tag">{word_count or "0"} 个逐词项目</span>
        </div>
      </a>"""
        )

    return '<div class="archive-list">\n      ' + "\n      ".join(cards) + "\n      </div>"


class WordbookEntryParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.entries: list[dict[str, str]] = []
        self._in_entry = False
        self._article_depth = 0
        self._in_h2 = False
        self._entry_text: list[str] = []
        self._entry_title: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attr_map = {name: value or "" for name, value in attrs}
        if tag == "article" and ("data-entry" in attr_map or attr_map.get("class") == "entry"):
            self._in_entry = True
            self._article_depth = 1
            self._entry_text = []
            self._entry_title = []
            return
        if not self._in_entry:
            return
        if tag == "article":
            self._article_depth += 1
        elif tag == "h2":
            self._in_h2 = True

    def handle_endtag(self, tag: str) -> None:
        if not self._in_entry:
            return
        if tag == "h2":
            self._in_h2 = False
        elif tag == "article":
            self._article_depth -= 1
            if self._article_depth <= 0:
                title = normalize_search_text(" ".join(self._entry_title))
                text = normalize_search_text(" ".join(self._entry_text))
                if text:
                    self.entries.append({"title": title or text[:48], "text": text})
                self._in_entry = False
                self._article_depth = 0

    def handle_data(self, data: str) -> None:
        if not self._in_entry:
            return
        text = data.strip()
        if not text:
            return
        self._entry_text.append(text)
        if self._in_h2:
            self._entry_title.append(text)


def normalize_search_text(value: str) -> str:
    return re.sub(r"\s+", " ", unescape(value_text(value))).strip()


def build_wordbook_search_index(wordbooks_dir: Path) -> dict[str, Any]:
    items: list[dict[str, Any]] = []
    if not wordbooks_dir.is_dir():
        return {"generated_at": datetime.now().isoformat(timespec="seconds"), "items": items}

    for child in sorted(wordbooks_dir.iterdir(), key=lambda path: path.name):
        if not child.is_dir():
            continue
        metadata = read_wordbook_metadata(child)
        index_path = child / "index.html"
        if not metadata or not index_path.is_file():
            continue
        parser = WordbookEntryParser()
        try:
            parser.feed(index_path.read_text(encoding="utf-8"))
        except OSError:
            continue
        folder = value_text(metadata.get("folder")) or child.name
        for entry in parser.entries:
            text = entry["text"]
            preview = text[:220] + ("..." if len(text) > 220 else "")
            items.append(
                {
                    "title": entry["title"],
                    "bookTitle": value_text(metadata.get("title")) or child.name,
                    "bookDate": value_text(metadata.get("generated_date")),
                    "sourceFile": value_text(metadata.get("source_file")),
                    "url": quote(folder) + "/",
                    "text": text,
                    "preview": preview,
                }
            )

    items.sort(key=lambda item: (value_text(item.get("bookDate")), value_text(item.get("bookTitle"))), reverse=True)
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "entry_count": len(items),
        "items": items,
    }


def rebuild_wordbook_search_index(wordbooks_dir: Path) -> Path:
    search_index_path = wordbooks_dir / "search-index.json"
    search_index_path.write_text(
        json.dumps(build_wordbook_search_index(wordbooks_dir), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return search_index_path


def rebuild_wordbook_index(wordbooks_dir: Path) -> Path:
    index_path = wordbooks_dir / "index.html"
    rebuild_wordbook_search_index(wordbooks_dir)
    archive_html = build_wordbook_archive_cards(wordbooks_dir)
    if not index_path.exists():
        index_path.write_text(build_wordbook_index_template(archive_html), encoding="utf-8")
        return index_path

    content = index_path.read_text(encoding="utf-8")
    start_marker = "<!-- WORDBOOK_ARCHIVE_START -->"
    end_marker = "<!-- WORDBOOK_ARCHIVE_END -->"
    start = content.find(start_marker)
    end = content.find(end_marker)
    if start == -1 or end == -1 or end <= start:
        index_path.write_text(build_wordbook_index_template(archive_html), encoding="utf-8")
        return index_path

    new_content = (
        content[: start + len(start_marker)]
        + "\n      "
        + archive_html
        + "\n      "
        + content[end:]
    )
    if "wordbook-search-section" not in new_content or "wordbook-search.js" not in new_content:
        new_content = build_wordbook_index_template(archive_html)
    index_path.write_text(new_content, encoding="utf-8")
    return index_path


def ensure_ielts_wordbook_entry(homepage_repo: Path) -> None:
    ielts_path = homepage_repo / "notes" / "ielts" / "index.html"
    if not ielts_path.is_file():
        return
    content = ielts_path.read_text(encoding="utf-8")
    if 'href="wordbooks/"' in content:
        return

    placeholder = """<h3>网页阅读版</h3>
        <p>这里预留给从 docx / pdf 整理出来的网页笔记正文。</p>
        <div class="note-slot">
          <p>暂无内容。可以先把原始文件放入 <strong>files/ielts/</strong>，再把重点内容整理到这里。</p>
        </div>"""
    replacement = """<h3>单词本收录</h3>
        <p>这里归档 Wgen 生成的手机阅读版 HTML，方便在电脑和 iPhone 上直接复习。</p>
        <div class="note-actions">
          <a class="button primary" href="wordbooks/">进入单词本收录</a>
        </div>"""
    if placeholder in content:
        ielts_path.write_text(content.replace(placeholder, replacement), encoding="utf-8")


def publish_mobile_html_to_homepage(
    *,
    data: dict[str, Any],
    source_path: Path,
    mobile_html_path: Path,
    homepage_repo: Path,
    push: bool,
    base_url: str = DEFAULT_HOMEPAGE_BASE_URL,
) -> HomepagePublishResult:
    homepage_repo = Path(homepage_repo)
    validate_homepage_repo(homepage_repo)

    wordbooks_dir = homepage_repo / "notes" / "ielts" / "wordbooks"
    tracked_targets = [
        "notes/ielts/index.html",
        "notes/ielts/wordbooks",
        "style.css",
    ]
    dirty_related = run_git(homepage_repo, ["status", "--porcelain", "--", *tracked_targets])
    if dirty_related.strip():
        raise ValueError(
            "个人主页仓库的雅思归档相关文件已有未提交改动，请先处理后再自动发布：\n"
            + dirty_related
        )

    wordbooks_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = f"{stamp}-{sanitize_url_part(source_path.stem)}"
    archive_dir = wordbooks_dir / folder
    archive_dir.mkdir(parents=True, exist_ok=False)

    shutil.copy2(mobile_html_path, archive_dir / "index.html")
    entry_count, word_count = count_wordbook_data(data)
    generated_date = datetime.now().strftime("%Y-%m-%d")
    metadata = {
        "title": f"{source_path.stem} 雅思单词本",
        "source_file": source_path.name,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "generated_date": generated_date,
        "entry_count": entry_count,
        "word_count": word_count,
    }
    (archive_dir / "meta.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    index_path = rebuild_wordbook_index(wordbooks_dir)
    ensure_ielts_wordbook_entry(homepage_repo)

    relative_paths = [
        "notes/ielts/index.html",
        "notes/ielts/wordbooks/index.html",
        "notes/ielts/wordbooks/search-index.json",
        f"notes/ielts/wordbooks/{folder}/index.html",
        f"notes/ielts/wordbooks/{folder}/meta.json",
    ]
    run_git(homepage_repo, ["add", "--", *relative_paths])
    staged = run_git(homepage_repo, ["diff", "--cached", "--name-only", "--", *relative_paths])
    commit_hash: str | None = None
    if staged.strip():
        message = f"Add IELTS wordbook: {source_path.stem}"
        run_git(homepage_repo, ["commit", "-m", message])
        if push:
            run_git(homepage_repo, ["push", "origin", "HEAD"])
        commit_hash = read_git_short_head(homepage_repo)

    url = f"{base_url.rstrip('/')}/notes/ielts/wordbooks/{quote(folder)}/"
    return HomepagePublishResult(
        archive_dir=archive_dir,
        index_path=index_path,
        commit_hash=commit_hash,
        url=url,
    )


def process_wordlist(
    *,
    input_docx: Path,
    output_dir: Path,
    rules_path: Path,
    api_key: str,
    model: str,
    provider_id: str = DEFAULT_PROVIDER_ID,
    thinking_enabled: bool = False,
    publish_to_homepage: bool = False,
    homepage_repo: Path | None = None,
    push_homepage: bool = True,
    progress_callback: ProgressCallback | None = None,
) -> ProcessingResult:
    input_docx = Path(input_docx)
    output_dir = Path(output_dir)
    rules_path = Path(rules_path)
    provider = get_api_provider(provider_id)

    if input_docx.suffix.lower() not in SUPPORTED_INPUT_SUFFIXES or not input_docx.is_file():
        raise FileNotFoundError("请选择一个有效的 .docx 或 .txt 原始单词记录文件。")
    if not api_key.strip():
        raise ValueError("API Key 不能为空。")
    if not model.strip():
        raise ValueError("模型名称不能为空。")
    if not rules_path.is_file():
        raise FileNotFoundError(f"找不到规则文件：{rules_path}")

    emit(progress_callback, 8, "正在读取原始单词记录……")
    source_text = extract_input_content(input_docx)
    rules = rules_path.read_text(encoding="utf-8").strip()
    if not rules:
        raise ValueError("规则文件为空。")

    emit(progress_callback, 22, f"已提取 {len(source_text):,} 个字符，准备调用 {provider.display_name}……")
    raw_text, usage = call_model(
        provider_id=provider.id,
        api_key=api_key.strip(),
        model=model.strip(),
        rules=rules,
        source_text=source_text,
        thinking_enabled=thinking_enabled,
    )

    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    debug_stem = sanitize_filename(input_docx.stem)
    raw_path = output_dir / f"{debug_stem}_{provider.raw_result_label}原始结果_{stamp}.json"
    raw_path.write_text(raw_text, encoding="utf-8")

    emit(progress_callback, 72, "正在校验模型返回的结构化数据……")
    data = parse_json_response(raw_text)
    validate_data(data)

    parsed_path = output_dir / f"{debug_stem}_结构化数据_{stamp}.json"
    parsed_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    emit(progress_callback, 84, "正在生成排版后的 Word 单词本……")
    output_docx = generate_word(data, input_docx, output_dir, stamp=stamp)
    emit(progress_callback, 94, "正在生成 iPhone 友好的手机版 HTML……")
    output_html = generate_mobile_html(data, input_docx, output_dir, stamp=stamp)

    homepage_publish = None
    if publish_to_homepage:
        if homepage_repo is None:
            raise ValueError("已启用个人主页归档，但没有设置个人主页仓库路径。")
        emit(progress_callback, 97, "正在归档到个人主页并推送 GitHub……")
        homepage_publish = publish_mobile_html_to_homepage(
            data=data,
            source_path=input_docx,
            mobile_html_path=output_html,
            homepage_repo=homepage_repo,
            push=push_homepage,
        )
    emit(progress_callback, 100, "整理完成。")

    prompt_tokens = getattr(usage, "prompt_tokens", None) if usage else None
    completion_tokens = getattr(usage, "completion_tokens", None) if usage else None
    total_tokens = getattr(usage, "total_tokens", None) if usage else None

    return ProcessingResult(
        output_docx=output_docx,
        output_html=output_html,
        raw_json=raw_path,
        parsed_json=parsed_path,
        model=model.strip(),
        homepage_publish=homepage_publish,
        prompt_tokens=prompt_tokens,
        completion_tokens=completion_tokens,
        total_tokens=total_tokens,
    )


def list_available_models(api_key: str, provider_id: str = DEFAULT_PROVIDER_ID) -> list[str]:
    provider = get_api_provider(provider_id)
    client = OpenAI(api_key=api_key.strip(), base_url=provider.base_url, timeout=60.0, max_retries=1)
    response = client.models.list()
    names = sorted({model.id for model in response.data if getattr(model, "id", None)})
    return names


def test_api(api_key: str, model: str, provider_id: str = DEFAULT_PROVIDER_ID) -> str:
    provider = get_api_provider(provider_id)
    client = OpenAI(api_key=api_key.strip(), base_url=provider.base_url, timeout=60.0, max_retries=1)
    request_args: dict[str, Any] = {
        "model": model.strip(),
        "messages": [{"role": "user", "content": "只回复：连接成功"}],
        "extra_body": build_extra_body(provider, False),
    }
    if provider.id in {"bigmodel", "zai"}:
        request_args["max_tokens"] = 32
    else:
        request_args["max_completion_tokens"] = 32
    response = client.chat.completions.create(**request_args)
    content = response.choices[0].message.content if response.choices else None
    return (content or "连接成功").strip()
